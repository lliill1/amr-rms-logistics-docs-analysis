"""Excel: каждый лист — таблица; длинные листы режутся на части с повтором строки заголовков."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .base import Block, ParsedDoc

ROWS_PER_BLOCK = 40
MAX_ROWS_PER_SHEET = 5000


def _fmt(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).replace("\n", " ").strip()


def parse_xlsx(path: Path, **_) -> ParsedDoc:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    for sheet_no, ws in enumerate(wb.worksheets, start=1):
        header: str | None = None
        buf: list[str] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS_PER_SHEET:
                break
            cells = [_fmt(v) for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            line = " | ".join(cells)
            if header is None:
                header = line
                continue
            buf.append(line)
            if len(buf) >= ROWS_PER_BLOCK:
                blocks.append(Block(text="\n".join([header, *buf]), kind="table", page=sheet_no, section=ws.title))
                buf = []
        if buf or (header and not buf):
            blocks.append(Block(text="\n".join([h for h in [header] if h] + buf), kind="table", page=sheet_no, section=ws.title))
    wb.close()
    return ParsedDoc(blocks=blocks, pages=len(wb.sheetnames), title=None, needs_ocr=False)
