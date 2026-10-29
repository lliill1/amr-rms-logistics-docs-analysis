"""Выгрузка отчёта в Word и таблиц в Excel. Оба файла полностью редактируемые."""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ROOT
from .db import rows

TEMPLATE = ROOT / "templates" / "report.docx"


def _src(r: dict) -> str:
    return f"{r['rel_path']}" + (f", с. {r['page']}" if r.get("page") else "")


def _requirements(conn, statuses: tuple[str, ...]) -> list[dict]:
    q = ",".join("?" * len(statuses))
    return rows(conn, f"""SELECT r.*, d.rel_path, d.title AS doc_title FROM requirements r
                          JOIN documents d ON d.id = r.doc_id WHERE r.status IN ({q})
                          ORDER BY r.target, r.category, r.code""", statuses)


def _documents(conn) -> list[dict]:
    return rows(conn, """SELECT d.*, s.rel_path AS superseded_by_path FROM documents d
                         LEFT JOIN documents s ON s.id = d.superseded_by ORDER BY d.relevance DESC, d.rel_path""")


def _table(doc, header: list[str], data: list[list], widths_cm: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = "" if v is None else str(v)
    if widths_cm:
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    return t


def build_docx(conn: sqlite3.Connection, out_dir: Path, include_auto: bool = False, project: str = "") -> Path:
    statuses = ("confirmed", "edited", "auto") if include_auto else ("confirmed", "edited")
    doc = Document(str(TEMPLATE)) if TEMPLATE.exists() else Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = max(sec.page_width, sec.page_height), min(sec.page_width, sec.page_height)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(2))

    docs = _documents(conn)
    reqs = _requirements(conn, statuses)
    questions = rows(conn, "SELECT * FROM open_questions WHERE status != 'closed' ORDER BY id")

    doc.add_heading("Роль AMR/RMS в логистических процессах предприятия", 0)
    doc.add_paragraph(f"Аналитический отчёт по архиву проектной документации{(' · ' + project) if project else ''}")
    doc.add_paragraph(f"Сформирован: {datetime.now():%d.%m.%Y %H:%M}. "
                      + ("Включены неподтверждённые факты (черновик)." if include_auto
                         else "В отчёт включены только факты, подтверждённые аналитиком."))

    # 1. Обзор архива
    doc.add_heading("1. Обзор архива и актуальные документы", 1)
    total = len(docs)
    parsed = sum(1 for d in docs if d["status"] == "parsed")
    relevant = [d for d in docs if d["relevance"] and d["relevance"] > 0 and d["is_current"]]
    outdated = [d for d in docs if not d["is_current"]]
    problems = [d for d in docs if d["status"] in ("needs_ocr", "unsupported", "error")]
    doc.add_paragraph(f"Всего файлов: {total}; с извлечённым текстом: {parsed}; актуальных и относящихся к AMR/RMS: "
                      f"{len(relevant)}; устаревших редакций: {len(outdated)}; требуют ручной обработки: {len(problems)}.")
    doc.add_heading("1.1. Документы, значимые для AMR/RMS", 2)
    _table(doc, ["Документ", "Тип", "Дата", "Ред.", "Статус", "Сторона", "Подсистемы", "Релев."],
           [[d["rel_path"], d["doc_type"], d["doc_date"], d["revision"], d["approval"], d["side"], d["subsystems"], d["relevance"]]
            for d in relevant], [8, 3.5, 2, 1.3, 2.2, 2.2, 3.5, 1.3])
    if outdated:
        doc.add_heading("1.2. Устаревшие редакции", 2)
        _table(doc, ["Документ", "Заменён документом"], [[d["rel_path"], d["superseded_by_path"] or "—"] for d in outdated], [12, 12])
    if problems:
        doc.add_heading("1.3. Документы, требующие ручной обработки", 2)
        reason = {"needs_ocr": "скан, нужен OCR", "unsupported": "формат не поддерживается", "error": "ошибка чтения"}
        _table(doc, ["Документ", "Причина"], [[d["rel_path"], reason.get(d["status"], d["status"]) + (f": {d['error']}" if d["error"] else "")]
                                              for d in problems], [12, 12])

    # 2–4. Разделы следующих версий
    for title in ("2. Общая схема логистики", "3. Карта процессов с участием AMR", "4. Схема взаимодействия систем"):
        doc.add_heading(title, 1)
        doc.add_paragraph("[Раздел будет формироваться автоматически в следующей версии сервиса. "
                          "Пока заполняется аналитиком вручную.]")

    # 5. Требования
    doc.add_heading("5. Сводная таблица функций и требований к AMR/RMS", 1)
    if not reqs:
        doc.add_paragraph("Подтверждённых требований пока нет.")
    for target in ("AMR", "RMS", "AMR+RMS", "Другое"):
        part = [r for r in reqs if r["target"] == target]
        if not part:
            continue
        doc.add_heading({"AMR": "Требования к роботам (AMR)", "RMS": "Требования к системе управления роботами (RMS)",
                         "AMR+RMS": "Общие требования к AMR и RMS", "Другое": "Требования к смежным системам и инфраструктуре"}[target], 2)
        _table(doc, ["ID", "Требование", "Категория", "Обязательность", "Источник", "Цитата"],
               [[r["code"], r["text"], r["category"], r["obligation"], _src(r), r["quote"]] for r in part],
               [1.6, 7, 2.5, 2.3, 5, 6.5])

    # 6. Вопросы
    doc.add_heading("6. Противоречия и открытые вопросы", 1)
    if not questions:
        doc.add_paragraph("Открытых вопросов пока нет.")
    else:
        _table(doc, ["№", "Вопрос", "Контекст", "Источники", "Статус"],
               [[q["id"], q["text"], q["context"], q["sources"], q["status"]] for q in questions], [1, 9, 6, 6, 2])

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"Отчёт_AMR_RMS_{datetime.now():%Y%m%d_%H%M}.docx"
    doc.save(str(path))
    return path


def build_xlsx(conn: sqlite3.Connection, out_dir: Path) -> Path:
    wb = Workbook()
    sheets = [
        ("Документы", ["ID", "Путь", "Название", "Тип", "Дата", "Редакция", "Статус документа", "Сторона", "Подсистемы",
                       "Релевантность", "Актуальный", "Состояние обработки"],
         [[d["id"], d["rel_path"], d["title"], d["doc_type"], d["doc_date"], d["revision"], d["approval"], d["side"],
           d["subsystems"], d["relevance"], "да" if d["is_current"] else "нет", d["status"]] for d in _documents(conn)]),
        ("Требования", ["ID", "Требование", "Адресат", "Категория", "Обязательность", "Источник", "Страница", "Цитата", "Статус ревью", "Комментарий"],
         [[r["code"], r["text"], r["target"], r["category"], r["obligation"], r["rel_path"], r["page"], r["quote"], r["status"], r["comment"]]
          for r in _requirements(conn, ("auto", "confirmed", "edited", "rejected"))]),
        ("Открытые вопросы", ["№", "Вопрос", "Контекст", "Источники", "Статус", "Ответ"],
         [[q["id"], q["text"], q["context"], q["sources"], q["status"], q["answer"]] for q in rows(conn, "SELECT * FROM open_questions ORDER BY id")]),
    ]
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="DCE6F5")
    for name, header, data in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in data[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"Таблицы_AMR_RMS_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(str(path))
    return path
