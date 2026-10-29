"""Сквозной тест на вымышленном демо-архиве: загрузка -> каталог -> поиск -> вопрос -> извлечение -> отчёт.

Запуск:  python -m pytest -q     (или: python -m unittest discover tests)
"""
from __future__ import annotations

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from core import catalog, db, report  # noqa: E402
from core.ask import ask  # noqa: E402
from core.config import load_settings  # noqa: E402
from core.extract import extract_requirements, quote_score  # noqa: E402
from core.index import make_chunks, search  # noqa: E402
from core.ingest import collect_files  # noqa: E402
from core.pipeline import run_all  # noqa: E402
from fake_llm import FakeLLM  # noqa: E402
from make_demo_archive import make  # noqa: E402


class PipelineTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        tmp = Path(cls.tmp.name)
        cls.archive = make(tmp / "archive")
        cls.settings = load_settings({"paths": {"data_dir": str(tmp / "data"), "output_dir": str(tmp / "out"),
                                                "db_file": str(tmp / "data" / "test.sqlite")}})
        cls.conn = db.connect(cls.settings.db_file)
        cls.llm = FakeLLM()
        cls.stats = run_all(cls.conn, cls.settings, cls.archive, llm=cls.llm)

    @classmethod
    def tearDownClass(cls):
        cls.conn.close()
        cls.tmp.cleanup()

    def doc(self, name_part: str) -> dict:
        r = db.one(self.conn, "SELECT * FROM documents WHERE rel_path LIKE ?", (f"%{name_part}%",))
        self.assertIsNotNone(r, name_part)
        return r

    def test_files_found_including_zip(self):
        self.assertEqual(self.stats.found, 5)
        self.assertEqual(self.stats.errors, 0)
        tz = self.doc("ТЗ_на_роботизацию")
        self.assertIn("Материалы_заказчика.zip/01_ТЗ/", tz["rel_path"])
        self.assertEqual(tz["status"], "parsed")

    def test_rerun_skips_duplicates(self):
        stats = run_all(self.conn, self.settings, self.archive, llm=self.llm)
        self.assertEqual(stats.added, 0)
        self.assertEqual(stats.duplicates, 5)

    def test_cards(self):
        v23 = self.doc("v2.3")
        self.assertEqual(v23["doc_type"], "Проектное решение")
        self.assertEqual(v23["revision"], "2.3")
        self.assertEqual(v23["doc_date"], "2025-09-20")
        self.assertEqual(v23["approval"], "утверждён")
        self.assertEqual(v23["side"], "подрядчик")
        self.assertGreaterEqual(v23["relevance"], 2)
        self.assertIn("EWM", v23["subsystems"])
        self.assertEqual(self.doc("v1.0")["approval"], "проект")
        self.assertEqual(self.doc("ТЗ_на")["doc_type"], "Техническое задание")
        self.assertEqual(self.doc("ТЗ_на")["side"], "заказчик")
        self.assertEqual(self.doc("Маппинг")["relevance"], 0)

    def test_revisions(self):
        old, new = self.doc("v1.0"), self.doc("v2.3")
        self.assertEqual(old["family_key"], new["family_key"])
        self.assertEqual(old["is_current"], 0)
        self.assertEqual(old["superseded_by"], new["id"])
        self.assertEqual(new["is_current"], 1)

    def test_search_prefers_current(self):
        hits = search(self.conn, "кто формирует транспортное задание", llm=self.llm, top_k=5)
        self.assertTrue(hits)
        self.assertTrue(all(h.is_current for h in hits))
        all_hits = search(self.conn, "транспортное задание MES", llm=self.llm, top_k=10, only_current=False)
        self.assertTrue(any("v1.0" in h.rel_path for h in all_hits))

    def test_ask(self):
        ans = ask(self.conn, self.settings, "Какая система создаёт транспортное задание для AMR?", llm=self.llm)
        self.assertTrue(ans.hits)
        self.assertIn("[1]", ans.text)
        self.assertIn("Фрагменты", self.llm.prompts[-1])
        no_llm = ask(self.conn, self.settings, "транспортное задание", llm=None)
        self.assertIsNone(no_llm.text)
        self.assertTrue(no_llm.hits)

    def test_extract_and_report(self):
        stats = extract_requirements(self.conn, self.settings, self.llm)
        self.assertGreater(stats["accepted"], 3)
        self.assertGreater(stats["rejected_quotes"], 0)  # выдуманное требование отброшено
        reqs = db.rows(self.conn, "SELECT * FROM requirements")
        self.assertFalse(any("летать" in r["text"] for r in reqs))
        self.assertTrue(all(r["code"].startswith("TR-") for r in reqs))
        # повторный запуск не дублирует
        again = extract_requirements(self.conn, self.settings, self.llm)
        self.assertEqual(again["chunks"], 0)

        db.update(self.conn, "requirements", reqs[0]["id"], status="confirmed")
        self.conn.execute("INSERT INTO open_questions(text, context) VALUES ('Кто инициирует задание: EWM или MES?', 'тест')")
        self.conn.commit()
        out = self.settings.output_dir
        docx_path = report.build_docx(self.conn, out)
        text = "\n".join(p.text for p in Document(str(docx_path)).paragraphs)
        self.assertIn("Сводная таблица функций и требований", text)
        tables = Document(str(docx_path)).tables
        all_cells = " ".join(c.text for t in tables for r in t.rows for c in r.cells)
        self.assertIn(reqs[0]["code"], all_cells)
        self.assertIn("EWM или MES", all_cells)
        xlsx_path = report.build_xlsx(self.conn, out)
        wb = load_workbook(str(xlsx_path))
        self.assertEqual(wb.sheetnames, ["Документы", "Требования", "Открытые вопросы"])
        self.assertEqual(wb["Требования"].max_row - 1, len(reqs))


class UnitTests(unittest.TestCase):
    def test_quote_score(self):
        src = "RMS должна принимать транспортные задания от EWM и возвращать статус выполнения."
        self.assertEqual(quote_score("RMS должна принимать транспортные задания от EWM", src), 1.0)
        self.assertEqual(quote_score("«RMS  должна принимать»", src), 1.0)
        self.assertLess(quote_score("Робот обязан летать над складом", src), 0.5)

    def test_dates_and_revisions(self):
        self.assertEqual(catalog.detect_date("Протокол от 12.03.2025.pdf", ""), "2025-03-12")
        self.assertEqual(catalog.detect_date("x.docx", "Москва «5» апреля 2024 г."), "2024-04-05")
        self.assertEqual(catalog.detect_revision("ПР_v2_1.docx", ""), "2.1")
        self.assertEqual(catalog.detect_revision("doc.docx", "Редакция 3"), "3")
        self.assertEqual(catalog.family_key("a/ПР_EWM_v1.0.docx"), catalog.family_key("b/ПР EWM v2.3 (1).docx"))

    def test_zip_cp866_names(self):
        with tempfile.TemporaryDirectory() as t:
            t = Path(t)
            zpath = t / "a.zip"
            name = "Папка/Документ.txt"
            class Cp866Info(zipfile.ZipInfo):
                # так пишут имена архиваторы русской Windows: cp866 без флага UTF-8
                def _encodeFilenameFlags(self):
                    return name.encode("cp866"), self.flag_bits & ~0x800

            with zipfile.ZipFile(zpath, "w") as zf:
                zf.writestr(Cp866Info(name), "робот")
            rep = collect_files(zpath, t / "work")
            self.assertEqual([f.rel_path for f in rep.files], ["a.zip/Папка/Документ.txt"])

    def test_chunks_respect_sections_and_size(self):
        blocks = [{"kind": "heading", "text": "1 A", "page": 1, "section": "1 A"},
                  {"kind": "text", "text": "x" * 1500, "page": 1, "section": "1 A"},
                  {"kind": "text", "text": "y" * 1500, "page": 2, "section": "1 A"},
                  {"kind": "heading", "text": "2 B", "page": 3, "section": "2 B"},
                  {"kind": "text", "text": "z" * 100, "page": 3, "section": "2 B"}]
        chunks = make_chunks(blocks, max_chars=2000, overlap=100)
        self.assertEqual([c["section"] for c in chunks], ["1 A", "1 A", "2 B"])
        self.assertTrue(all(len(c["text"]) <= 2200 for c in chunks))
        self.assertEqual((chunks[1]["page_from"], chunks[1]["page_to"]), (1, 2))


if __name__ == "__main__":
    unittest.main()
