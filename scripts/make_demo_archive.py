"""Создаёт ВЫМЫШЛЕННЫЙ демо-архив для проверки сервиса без реальных документов заказчика.

    python scripts/make_demo_archive.py            # -> data/demo_archive
    python scripts/make_demo_archive.py D:/tmp/demo

Внутри: две редакции проектного решения (с противоречием), ТЗ в zip-архиве с русскими именами,
таблица маппинга Excel, протокол совещания (PDF, если установлен reportlab, иначе .txt).
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent


def _docx(path: Path, title: str, head: list[str], sections: list[tuple[str, list[str]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.core_properties.title = title
    for line in head:
        d.add_paragraph(line)
    d.add_heading(title, 0)
    for heading, paras in sections:
        d.add_heading(heading, 1)
        for p in paras:
            d.add_paragraph(p)
    d.save(str(path))


def make(dest: Path) -> Path:
    dest.mkdir(parents=True, exist_ok=True)

    _docx(dest / "02_EWM" / "ПР_EWM_Внутрипроизводственная_логистика_v1.0.docx",
          "Проектное решение. Внутрипроизводственная логистика SAP EWM",
          ["Редакция 1.0", "Дата: 15.02.2024", "Проект (на согласовании)", "Исполнитель: подрядчик"],
          [("1 Общие положения", ["Документ описывает процессы внутрипроизводственной логистики после внедрения SAP EWM."]),
           ("2 Транспортные задания для AMR",
            ["Транспортное задание на подвоз комплектующих к участку сборки формирует MES по сигналу оператора.",
             "Система управления роботами RMS получает задание от MES и назначает свободного робота AMR."])])

    _docx(dest / "02_EWM" / "ПР_EWM_Внутрипроизводственная_логистика_v2.3.docx",
          "Проектное решение. Внутрипроизводственная логистика SAP EWM",
          ["Редакция 2.3", "Дата: 20.09.2025", "УТВЕРЖДАЮ", "Исполнитель: подрядчик"],
          [("1 Общие положения", ["Документ описывает процессы внутрипроизводственной логистики после внедрения SAP EWM."]),
           ("2 Транспортные задания для AMR",
            ["Транспортное задание формирует SAP EWM при создании складской задачи на перемещение в зону подачи PSA.",
             "RMS должна принимать транспортные задания от EWM и возвращать статус выполнения каждого задания.",
             "Подтверждение доставки фиксируется в EWM после сигнала RMS о завершении задания."]),
           ("3 Требования к мобильным роботам",
            ["Робот AMR должен перевозить тележки массой до 600 кг.",
             "Робот должен автоматически открывать скоростные ворота между корпусами 1 и 2 через интеграцию RMS с контроллером ворот.",
             "Желательно, чтобы время работы робота от одной зарядки составляло не менее 8 часов."])])

    tz = dest / "_tmp" / "ТЗ_на_роботизацию_внутрицеховой_логистики.docx"
    _docx(tz, "Техническое задание на роботизацию внутрицеховой логистики",
          ["Заказчик: [предприятие]", "Дата: 01.11.2023"],
          [("1 Назначение", ["Роботизация подвоза материалов со склада к участкам сборки с помощью мобильных роботов."]),
           ("2 Требования к системе управления роботами",
            ["Система управления роботами должна поддерживать приоритеты транспортных заданий.",
             "RMS должна обеспечивать обход препятствий и перестроение маршрута без участия оператора."])])
    with zipfile.ZipFile(dest / "Материалы_заказчика.zip", "w") as zf:
        zf.write(tz, "01_ТЗ/" + tz.name)
    tz.unlink()
    tz.parent.rmdir()

    wb = Workbook()
    ws = wb.active
    ws.title = "Маппинг"
    ws.append(["Поле ERP", "Поле EWM", "Описание"])
    ws.append(["MATNR", "PRODUCT", "Материал"])
    ws.append(["AUFNR", "DOCNO", "Производственный заказ"])
    (dest / "04_Интеграция").mkdir(parents=True, exist_ok=True)
    wb.save(str(dest / "04_Интеграция" / "Маппинг_ERP_EWM.xlsx"))

    protocol_lines = [
        "Протокол рабочего совещания по интеграции RMS от 12.03.2025",
        "Присутствовали: представители заказчика и подрядчика.",
        "1. Обсудили, какая система инициирует транспортное задание для AMR: EWM или MES.",
        "Решение не принято, вопрос вынесен на следующее совещание.",
        "2. Протокол обмена между RMS и роботами: рассматривается VDA 5050.",
    ]
    (dest / "05_Протоколы").mkdir(parents=True, exist_ok=True)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas

        font = next((p for p in [Path("C:/Windows/Fonts/arial.ttf"),
                                 Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")] if p.exists()), None)
        if font is None:
            raise ImportError("нет шрифта с кириллицей")
        pdfmetrics.registerFont(TTFont("Cyr", str(font)))
        c = canvas.Canvas(str(dest / "05_Протоколы" / "Протокол_совещания_RMS_от_12.03.2025.pdf"), pagesize=A4)
        c.setFont("Cyr", 11)
        y = 800
        for line in protocol_lines:
            c.drawString(50, y, line)
            y -= 20
        c.save()
    except ImportError:
        (dest / "05_Протоколы" / "Протокол_совещания_RMS_от_12.03.2025.txt").write_text("\n\n".join(protocol_lines), encoding="utf-8")
    return dest


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "demo_archive"
    print("Демо-архив создан:", make(target))
